#!/usr/bin/env python3
"""拉取上游 CPS 广告收益接口（POST /open/order/adSummary）逐日数据，原样导出 xlsx（离线查看用）。

用法：
    python3 fetch_ad_summary_to_xlsx.py                # 默认 2026-08-01 ~ 今天
    python3 fetch_ad_summary_to_xlsx.py --start 2026-10-01

接口实测行为（2026-10-14 探测）：
- 鉴权与 order/list 相同：先 POST /open/oauth2/token（HTTP Basic）换 Bearer
- 单次查询区间 ≤ 1 天 → 逐日查询（当日 00:00:00 – 23:59:59，UTC）
- 响应行字段：date（日期）、ad_revenue_dollars（**美元单位**，不是最小单位！）
- ⚠ 2026-10-14 现状：接口只返回最近约 2 天数据，更早日期返回空 list（xlsx 里留空）
- ⚠ 当天数据持续累加中，只有次日才定稿

凭据：默认读 partner-dashboard/backend/.env 的 UPSTREAM_CLIENT_ID / UPSTREAM_CLIENT_SECRET。
"""

from __future__ import annotations

import argparse
import os
import sys
import time
from datetime import date, datetime, timedelta
from pathlib import Path

import requests
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

DEFAULT_ENV_FILE = Path(__file__).resolve().parents[1] / "backend" / ".env"
DEFAULT_BASE = "https://adminapi.netshort.com/prod-api/cps"
DEFAULT_START = date(2026, 8, 1)

PATH_TOKEN = "/open/oauth2/token"
PATH_AD_SUMMARY = "/open/order/adSummary"

HTTP_ATTEMPTS = 3
HTTP_RETRY_BACKOFF = 2.0
REQUEST_TIMEOUT = 30


class ScriptError(Exception):
    pass


def load_env(env_file: Path) -> dict[str, str]:
    env: dict[str, str] = {}
    if not env_file.exists():
        return env
    for raw in env_file.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, _, value = line.partition("=")
        env[key.strip()] = value.strip().strip("'\"")
    return env


def pick(env: dict[str, str], key: str, default: str = "") -> str:
    return os.environ.get(key, "") or env.get(key, "") or default


class CpsClient:
    def __init__(self, base: str, client_id: str, client_secret: str) -> None:
        self.base = base.rstrip("/")
        self.client_id = client_id
        self.client_secret = client_secret
        self._token: str | None = None
        self._token_expires_at = float(0)
        self.request_count = 0

    def _fetch_token(self) -> str:
        try:
            resp = requests.post(
                self.base + PATH_TOKEN,
                auth=(self.client_id, self.client_secret),
                data={"grant_type": "client_credentials"},
                timeout=REQUEST_TIMEOUT,
            )
        except requests.RequestException as exc:
            raise ScriptError(f"换票请求失败：{type(exc).__name__}") from exc
        if resp.status_code == 401:
            raise ScriptError("换票被拒（HTTP 401）：client_id / client_secret 不正确")
        try:
            payload = resp.json()
        except ValueError as exc:
            raise ScriptError("换票响应不是 JSON") from exc
        data = payload.get("data") if isinstance(payload, dict) else None
        token = data.get("accessToken") or data.get("access_token") if isinstance(data, dict) else None
        if not isinstance(payload, dict) or payload.get("code") != 200 or not token:
            msg = payload.get("msg") if isinstance(payload, dict) else resp.text[:200]
            raise ScriptError(f"换票失败：{msg}")
        try:
            ttl = float(data.get("expiresIn") if data.get("expiresIn") is not None else data.get("expires_in"))
        except (TypeError, ValueError):
            ttl = 3600.0
        self._token = str(token)
        self._token_expires_at = time.time() + (ttl if ttl > 0 else 3600.0)
        return self._token

    def get_token(self, *, force: bool = False) -> str:
        if not force and self._token and time.time() < self._token_expires_at - 60:
            return self._token
        return self._fetch_token()

    def fetch_day(self, day: date) -> float | None:
        """拉单日广告收益（美元）。空 list → None（接口无该日数据）；失败 → 抛 ScriptError。"""
        body = {
            "startDate": f"{day.isoformat()} 00:00:00",
            "endDate": f"{day.isoformat()} 23:59:59",
            "page": 1,
            "pageSize": 100,
        }
        refreshed = False
        last_err, attempt = "", 0
        for attempt in range(1, HTTP_ATTEMPTS + 1):
            try:
                resp = requests.post(
                    self.base + PATH_AD_SUMMARY,
                    json=body,
                    headers={"Authorization": f"Bearer {self.get_token()}"},
                    timeout=REQUEST_TIMEOUT,
                )
            except requests.RequestException as exc:
                last_err = f"transport:{type(exc).__name__}"
            else:
                self.request_count += 1
                if resp.status_code == 401 and not refreshed:
                    refreshed = True
                    self.get_token(force=True)
                    continue
                if resp.status_code != 200:
                    last_err = f"http:{resp.status_code}"
                else:
                    try:
                        payload = resp.json()
                    except ValueError:
                        last_err = "not-json"
                    else:
                        if isinstance(payload, dict) and payload.get("code") == 401 and not refreshed:
                            refreshed = True
                            self.get_token(force=True)
                            continue
                        if not isinstance(payload, dict) or payload.get("code") != 200:
                            msg = payload.get("msg") if isinstance(payload, dict) else ""
                            trace = payload.get("traceId") if isinstance(payload, dict) else ""
                            raise ScriptError(f"code={payload.get('code')} {msg} (traceId={trace})")
                        data = payload.get("data")
                        rows = data.get("list") if isinstance(data, dict) else None
                        if not isinstance(rows, list) or not rows:
                            return None
                        for row in rows:
                            if str(row.get("date") or "")[:10] == day.isoformat():
                                return float(row["ad_revenue_dollars"])
                        return None
            if last_err.startswith("transport:") or last_err == "http:429" or last_err.startswith("http:5"):
                if attempt < HTTP_ATTEMPTS:
                    time.sleep(HTTP_RETRY_BACKOFF * attempt)
                    continue
            break
        raise ScriptError(f"{PATH_AD_SUMMARY} 请求失败（尝试 {attempt} 次）：{last_err}")


def main() -> int:
    parser = argparse.ArgumentParser(description="拉取 adSummary 逐日广告收益并导出 xlsx（纯接口）")
    parser.add_argument("--start", help=f"起始日期（默认 {DEFAULT_START}）")
    parser.add_argument("--end", help="结束日期（默认今天）")
    parser.add_argument("--out", help="输出 xlsx 路径（默认 ad_summary_export_<时间戳>.xlsx）")
    parser.add_argument("--env", default=str(DEFAULT_ENV_FILE))
    parser.add_argument("--base", help="上游网关地址")
    args = parser.parse_args()

    env = load_env(Path(args.env))
    client_id = pick(env, "UPSTREAM_CLIENT_ID")
    client_secret = pick(env, "UPSTREAM_CLIENT_SECRET")
    if not client_id or not client_secret:
        print(f"错误：未找到上游凭据（{args.env}）", file=sys.stderr)
        return 2
    base = args.base or pick(env, "UPSTREAM_BASE", DEFAULT_BASE)

    end_day = date.fromisoformat(args.end) if args.end else date.today()
    start_day = date.fromisoformat(args.start) if args.start else DEFAULT_START
    out_path = Path(args.out).resolve() if args.out else Path(
        f"ad_summary_export_{datetime.now():%Y%m%d_%H%M%S}.xlsx").resolve()

    print(f"上游：{base}")
    print(f"范围：{start_day} ~ {end_day}（UTC 日界，逐日查询），输出：{out_path}")

    client = CpsClient(base, client_id, client_secret)
    api_data: dict[str, float] = {}
    empty_days: list[date] = []
    failed_days: list[date] = []
    day = start_day
    total_days = (end_day - start_day).days + 1
    index = 0
    while day <= end_day:
        index += 1
        try:
            value = client.fetch_day(day)
        except ScriptError as exc:
            print(f"\n✗ {day} 失败：{exc}", file=sys.stderr)
            failed_days.append(day)
        else:
            if value is None:
                empty_days.append(day)
            else:
                api_data[day.isoformat()] = value
        print(f"\r[{index}/{total_days}] {day} 有值天数={len(api_data)}", end="", flush=True)
        day += timedelta(days=1)
    print()

    wb = Workbook()
    ws = wb.active
    ws.title = "ad_daily"
    ws.append(["date", "ad_revenue_usd"])
    day = start_day
    while day <= end_day:
        ws.append([day.isoformat(), api_data.get(day.isoformat(), "")])
        day += timedelta(days=1)
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 18
    ws.freeze_panes = "A2"

    total = sum(api_data.values())
    info = wb.create_sheet("导出说明")
    meta = [
        ("生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("接口", "POST /open/order/adSummary（逐日查询，UTC 日界，纯接口数据）"),
        ("单位", "ad_revenue_dollars 为美元（注意：不是最小单位/美分）"),
        ("有值天数 / 合计", f"{len(api_data)} 天 / {total:,.2f} USD"),
        ("接口返回空的日期数", str(len(empty_days))),
        ("失败日期", ", ".join(d.isoformat() for d in failed_days) or "无"),
        ("⚠ 调查发现", "2026-10-14 实测：接口只返回最近约 2 天数据，更早日期一律空 list"),
        ("备注", "当天数据持续累加，次日才定稿；本表只含接口当前能给的数据"),
    ]
    for k, v in meta:
        info.append([k, v])
    info.column_dimensions["A"].width = 24
    info.column_dimensions["B"].width = 90

    wb.save(out_path)
    print(f"有值: {len(api_data)} 天，合计 {total:,.2f} USD；空: {len(empty_days)} 天")
    print(f"已导出：{out_path}")
    return 0 if not failed_days else 1


if __name__ == "__main__":
    sys.exit(main())
