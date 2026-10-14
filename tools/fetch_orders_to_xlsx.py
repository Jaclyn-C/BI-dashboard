#!/usr/bin/env python3
"""拉取上游 CPS 订单接口（POST /open/order/list）的全部数据，原样存入 xlsx（离线查看用）。

用法：
    python3 fetch_orders_to_xlsx.py                    # 默认 2026-08-01 ~ 今天
    python3 fetch_orders_to_xlsx.py --start 2026-09-01
    python3 fetch_orders_to_xlsx.py --out 我的订单.xlsx

接口实测行为（2026-10-14 探测）：
- 鉴权：POST /open/oauth2/token（HTTP Basic，client_id:client_secret）换 Bearer token
- 单次查询区间 ≤ 1 天（超 1 天返回 code=400「查询时间范围不能超过1天」）→ 按天迭代
- pageSize 必须在 [100, 1000]
- pageSize=1000 时单日一页返回全部（hasMore=False）；pageSize=100 时翻页是坏的
  （不同 page 返回相同数据）→ 固定用 1000，并对结果按 orderId 全局去重防御
- 响应信封 {code, msg, data:{list, page, pageSize, hasMore, lastId, timeZero}, traceId}
- 订单字段（12 个）：orderId, theaterId, chapterId, orderAmount, currencyType, amount,
  createTime, refundTime, promotionCode, linkId, language, orderStatus
  （createTime 为 UTC，接口返回 timeZero=+00:00；当天数据有同步延迟，不全）
- 数据自 2026-08-05 起（此前为空），默认起始 2026-08-01 已覆盖全部

凭据：默认读 partner-dashboard/backend/.env 的 UPSTREAM_CLIENT_ID / UPSTREAM_CLIENT_SECRET
（可选 UPSTREAM_BASE），同名环境变量优先。脚本不打印、不存储凭据。
"""

from __future__ import annotations

import argparse
import os
import sys
import time
from datetime import date, datetime, timedelta
from pathlib import Path

import requests
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

DEFAULT_ENV_FILE = Path(__file__).resolve().parents[1] / "backend" / ".env"
DEFAULT_BASE = "https://adminapi.netshort.com/prod-api/cps"
DEFAULT_START = date(2026, 8, 1)  # 接口数据自 2026-08-05 起，此前全空

PATH_TOKEN = "/open/oauth2/token"
PATH_ORDER_LIST = "/open/order/list"

PAGE_SIZE = 1000        # 接口上限；pageSize=100 时翻页会返回重复数据，勿改小
MAX_PAGES_PER_DAY = 20  # hasMore 翻页保险丝（正常单日 1 页即完）
HTTP_ATTEMPTS = 3
HTTP_RETRY_BACKOFF = 2.0
REQUEST_TIMEOUT = 60

FIELD_NOTES = [
    ("orderId", "订单号"),
    ("theaterId", "剧场/短剧 ID（可能为空）"),
    ("chapterId", "章节 ID（可能为空）"),
    ("orderAmount", "美元金额（最小单位=美分，上游已折算）"),
    ("currencyType", "原币币种，如 USD / BRL / IDR / JPY"),
    ("amount", "原币金额（最小单位）"),
    ("createTime", "下单时间（UTC，接口 timeZero=+00:00；当天数据有延迟不全）"),
    ("refundTime", "退款时间；非空 = 已退款"),
    ("promotionCode", "渠道推广码"),
    ("linkId", "推广链接 ID"),
    ("language", "语言（如 en_US / pt_PT，可能为空）"),
    ("orderStatus", "订单状态（整数，实测均为 2）"),
]


class ScriptError(Exception):
    """致命错误（凭据缺失、信封错误、重试耗尽等）。"""


# ---------- .env ----------

def load_env(env_file: Path) -> dict[str, str]:
    env: dict[str, str] = {}
    if not env_file.exists():
        return env
    for raw in env_file.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, _, value = line.partition("=")
        env[key.strip()] = value.strip().strip("'\"")
    return env


def pick(env: dict[str, str], key: str, default: str = "") -> str:
    return os.environ.get(key, "") or env.get(key, "") or default


# ---------- 接口客户端 ----------

class CpsClient:
    def __init__(self, base: str, client_id: str, client_secret: str) -> None:
        self.base = base.rstrip("/")
        self.client_id = client_id
        self.client_secret = client_secret
        self._token: str | None = None
        self._token_expires_at: float = 0.0
        self.request_count = 0

    def _fetch_token(self) -> str:
        try:
            resp = requests.post(
                self.base + PATH_TOKEN,
                auth=(self.client_id, self.client_secret),
                data={"grant_type": "client_credentials"},
                timeout=REQUEST_TIMEOUT,
            )
        except requests.RequestException as exc:
            raise ScriptError(f"换票请求失败：{type(exc).__name__}") from exc
        if resp.status_code == 401:
            raise ScriptError("换票被拒（HTTP 401）：client_id / client_secret 不正确")
        try:
            payload = resp.json()
        except ValueError as exc:
            raise ScriptError("换票响应不是 JSON") from exc
        data = payload.get("data") if isinstance(payload, dict) else None
        token = data.get("accessToken") or data.get("access_token") if isinstance(data, dict) else None
        if not isinstance(payload, dict) or payload.get("code") != 200 or not token:
            msg = payload.get("msg") if isinstance(payload, dict) else resp.text[:200]
            raise ScriptError(f"换票失败：{msg}")
        try:
            ttl = float(data.get("expiresIn") if data.get("expiresIn") is not None else data.get("expires_in"))
        except (TypeError, ValueError):
            ttl = 3600.0
        self._token = str(token)
        self._token_expires_at = time.time() + (ttl if ttl > 0 else 3600.0)
        return self._token

    def get_token(self, *, force: bool = False) -> str:
        if not force and self._token and time.time() < self._token_expires_at - 60:
            return self._token
        return self._fetch_token()

    def _post_once(self, body: dict, token: str):
        try:
            resp = requests.post(
                self.base + PATH_ORDER_LIST,
                json=body,
                headers={"Authorization": f"Bearer {token}"},
                timeout=REQUEST_TIMEOUT,
            )
        except requests.RequestException as exc:
            return None, f"transport:{type(exc).__name__}"
        self.request_count += 1
        if resp.status_code != 200:
            return None, f"http:{resp.status_code}"
        try:
            return resp.json(), None
        except ValueError:
            return None, "not-json"

    def _request(self, body: dict):
        """单次业务调用：网络/5xx/429 退避重试；401（HTTP 层）重换票一次。"""
        refreshed = False
        last_err, attempt = "", 0
        for attempt in range(1, HTTP_ATTEMPTS + 1):
            payload, err = self._post_once(body, self.get_token())
            if err is None:
                return payload
            last_err = err
            if err == "http:401" and not refreshed:
                refreshed = True
                self.get_token(force=True)
                continue
            retryable = err.startswith("transport:") or err == "http:429" or err.startswith("http:5")
            if not retryable or attempt == HTTP_ATTEMPTS:
                break
            time.sleep(HTTP_RETRY_BACKOFF * attempt)
        raise ScriptError(f"{PATH_ORDER_LIST} 请求失败（尝试 {attempt} 次）：{last_err}")

    def fetch_day(self, day: date) -> list[dict]:
        """拉单个自然日（UTC 日界）订单：page=1 起，按 hasMore 翻页。"""
        rows: list[dict] = []
        for page in range(1, MAX_PAGES_PER_DAY + 1):
            body = {
                "startDate": f"{day.isoformat()} 00:00:00",
                "endDate": f"{day.isoformat()} 23:59:59",
                "page": page,
                "pageSize": PAGE_SIZE,
            }
            payload = self._request(body)
            if isinstance(payload, dict) and payload.get("code") == 401:
                self.get_token(force=True)  # 信封层 401：重换票重放一次
                payload = self._request(body)
            if not isinstance(payload, dict) or payload.get("code") != 200:
                msg = payload.get("msg") if isinstance(payload, dict) else ""
                trace = payload.get("traceId") if isinstance(payload, dict) else ""
                raise ScriptError(f"code={payload.get('code') if isinstance(payload, dict) else '?'} {msg} (traceId={trace})")
            data = payload.get("data")
            if isinstance(data, list):  # 实测偶发 data 直接是数组，防御处理
                rows.extend(r for r in data if isinstance(r, dict))
                break
            if not isinstance(data, dict):
                break
            page_rows = [r for r in (data.get("list") or []) if isinstance(r, dict)]
            rows.extend(page_rows)
            if not data.get("hasMore") or not page_rows:
                break
        return rows


# ---------- xlsx ----------

def to_cell(value):
    """原样写入；dict/list 转字符串防 openpyxl 报错。"""
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        import json
        return json.dumps(value, ensure_ascii=False)
    return value


def write_xlsx(path: Path, rows: list[dict], meta: list[tuple[str, str]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "orders"

    # 列 = 第一行字段的自然顺序（接口返回顺序），后续行的新字段追加
    columns: list[str] = []
    for row in rows:
        for key in row:
            if key not in columns:
                columns.append(key)

    ws.append(columns)
    for row in rows:
        ws.append([to_cell(row.get(c)) for c in columns])
    for idx, name in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = max(12, min(30, len(name) * 2 + 4))
    ws.freeze_panes = "A2"

    info = wb.create_sheet("导出说明")
    for key, value in meta:
        info.append([key, value])
    info.append([])
    info.append(["字段", "语义（实测）"])
    for name, note in FIELD_NOTES:
        info.append([name, note])
    for extra in [c for c in columns if c not in {n for n, _ in FIELD_NOTES}]:
        info.append([extra, "（新出现的字段，语义未知）"])
    info.column_dimensions["A"].width = 24
    info.column_dimensions["B"].width = 90

    wb.save(path)


# ---------- 主流程 ----------

def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="拉取上游 CPS order/list 全部数据并导出 xlsx")
    parser.add_argument("--start", help=f"起始日期 YYYY-MM-DD（默认 {DEFAULT_START}，数据最早 2026-08-05）")
    parser.add_argument("--end", help="结束日期 YYYY-MM-DD（默认今天）")
    parser.add_argument("--out", help="输出 xlsx 路径（默认 orders_export_<时间戳>.xlsx）")
    parser.add_argument("--env", default=str(DEFAULT_ENV_FILE), help=".env 路径")
    parser.add_argument("--base", help="上游网关地址（默认 .env 的 UPSTREAM_BASE 或商用地址）")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    env = load_env(Path(args.env))

    client_id = pick(env, "UPSTREAM_CLIENT_ID")
    client_secret = pick(env, "UPSTREAM_CLIENT_SECRET")
    if not client_id or not client_secret:
        print(f"错误：未找到上游凭据。请在 {args.env} 配置 UPSTREAM_CLIENT_ID / UPSTREAM_CLIENT_SECRET。", file=sys.stderr)
        return 2
    base = args.base or pick(env, "UPSTREAM_BASE", DEFAULT_BASE)

    end_day = date.fromisoformat(args.end) if args.end else date.today()
    start_day = date.fromisoformat(args.start) if args.start else DEFAULT_START
    if start_day > end_day:
        print("错误：--start 晚于 --end", file=sys.stderr)
        return 2
    out_path = Path(args.out).resolve() if args.out else Path(f"orders_export_{datetime.now():%Y%m%d_%H%M%S}.xlsx").resolve()

    total_days = (end_day - start_day).days + 1
    print(f"上游：{base}")
    print(f"范围：{start_day} ~ {end_day}（{total_days} 天，UTC 日界，逐日单次查询）")
    print(f"输出：{out_path}")

    client = CpsClient(base, client_id, client_secret)
    all_rows: list[dict] = []
    seen: set[str] = set()
    duplicates = 0
    full_page_days: list[date] = []
    failed_days: list[date] = []
    started = time.time()
    aborted = False

    try:
        day, index = start_day, 0
        while day <= end_day:
            index += 1
            try:
                rows = client.fetch_day(day)
            except ScriptError as exc:
                print(f"\n✗ {day} 失败：{exc}", file=sys.stderr)
                failed_days.append(day)
            else:
                if len(rows) >= PAGE_SIZE:
                    full_page_days.append(day)  # 单日超一页，翻页接口不可靠，可能不完整
                for row in rows:
                    key = str(row.get("orderId"))
                    if key in seen:
                        duplicates += 1
                        continue
                    seen.add(key)
                    all_rows.append(row)
            print(f"\r[{index}/{total_days}] {day} 累计 {len(all_rows)} 单", end="", flush=True)
            if index % 30 == 0:
                print()
            day += timedelta(days=1)
    except KeyboardInterrupt:
        aborted = True
        print("\n收到中断，导出已拉到的部分……")

    print()
    print(f"完成：{len(all_rows)} 单（去重丢弃 {duplicates} 条），{client.request_count} 次请求，"
          f"耗时 {time.time() - started:.0f}s")
    if full_page_days:
        print(f"⚠ 以下日期返回满页（≥{PAGE_SIZE} 行），因翻页不可靠可能有截断："
              f"{', '.join(d.isoformat() for d in full_page_days)}")
    if failed_days:
        print(f"✗ 拉取失败（缺数据）：{', '.join(d.isoformat() for d in failed_days)}")
    if aborted:
        print("⚠ 本次导出不完整（用户中断）")

    meta = [
        ("生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("上游网关", base),
        ("接口", "POST /open/order/list（逐日查询，UTC 日界，pageSize=1000）"),
        ("时间范围", f"{start_day} ~ {end_day}"),
        ("总行数（按 orderId 去重后）", str(len(all_rows))),
        ("重复丢弃", str(duplicates)),
        ("业务请求数", str(client.request_count)),
        ("返回满页的日期（可能有截断）", ", ".join(d.isoformat() for d in full_page_days) or "无"),
        ("拉取失败的日期", ", ".join(d.isoformat() for d in failed_days) or "无"),
        ("完整性", "不完整（用户中断）" if aborted else "完整"),
        ("备注", "数据为接口原始字段原样导出；createTime 为 UTC，当天数据有同步延迟不全"),
    ]
    write_xlsx(out_path, all_rows, meta)
    print(f"已导出：{out_path}")
    return 0 if not (failed_days or aborted) else 1


if __name__ == "__main__":
    sys.exit(main())
